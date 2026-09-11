"""Second-order boundary finite-difference solver for Problem 1.

Unknowns are the grid nodes 0 <= r < R.  The surface value is eliminated
with a second-order one-sided Robin condition, which avoids storing a
spurious half control volume at the boundary.
"""
from __future__ import annotations
import numpy as np
from pathlib import Path
from openpyxl import load_workbook
from scipy.sparse import eye, lil_matrix
from scipy.sparse.linalg import spsolve

ROOT = Path(__file__).resolve().parents[1]


def read_air():
    wb = load_workbook(ROOT / "附件" / "附件1.xlsx", read_only=True, data_only=True)
    a = np.asarray(list(wb.active.iter_rows(min_row=2, values_only=True)), dtype=float)
    wb.close()
    return a[:, 0], a[:, 1], a[:, 2]


def harmonic(a, b):
    return 2.0 * a * b / (a + b)


def boundary_coeff(c_surface: float, robin_h: float, air: float, dr: float):
    den = 3.0 * c_surface + 2.0 * dr * robin_h
    return 4.0 * c_surface / den, -c_surface / den, 2.0 * dr * robin_h * air / den


def assemble(r: np.ndarray, coeff_full: np.ndarray, robin_h: float, air: float, capacity: float):
    """Assemble q_t=Aq+b for unknowns q[0:N] (surface q[N] eliminated)."""
    n = len(r) - 1
    dr = r[1] - r[0]
    A = lil_matrix((n, n), dtype=float)
    b = np.zeros(n, dtype=float)
    rf = r[:-1] + 0.5 * dr
    cf = harmonic(coeff_full[:-1], coeff_full[1:])

    # center: cylindrical symmetry limit
    c = 4.0 * coeff_full[0] / (dr * dr * capacity)
    A[0, 0] = -c
    A[0, 1] = c

    # interior nodes
    for i in range(1, n - 1):
        cm = rf[i - 1] * cf[i - 1] / (r[i] * dr * dr * capacity)
        cp = rf[i] * cf[i] / (r[i] * dr * dr * capacity)
        A[i, i - 1] = cm
        A[i, i] = -(cm + cp)
        A[i, i + 1] = cp

    # node n-1, with q[n] = a1*q[n-1] + a2*q[n-2] + ba
    a1, a2, ba = boundary_coeff(coeff_full[-1], robin_h, air, dr)
    cm = rf[-2] * cf[-2] / (r[-2] * dr * dr * capacity)
    cp = rf[-1] * cf[-1] / (r[-2] * dr * dr * capacity)
    A[n - 1, n - 2] = cp * a2 + cm
    A[n - 1, n - 1] = cp * (a1 - 1.0) - cm
    b[n - 1] = cp * ba
    return A, b, (a1, a2, ba)


def surface_from_inner(q_inner: np.ndarray, coeff_surface: float, robin_h: float, air: float, dr: float):
    a1, a2, ba = boundary_coeff(coeff_surface, robin_h, air, dr)
    return a1 * q_inner[-1] + a2 * q_inner[-2] + ba


def solve_q1(dr_cm=0.05, dt=0.5, t_end=1800.0):
    at, aa, ac = read_air()
    dr = dr_cm / 100.0
    n = int(round(0.02 / dr))
    r = np.arange(n + 1, dtype=float) * dr
    rho, cp, k, h, hm = 820.0, 2600.0, 0.36, 25.0, 8e-7
    cap = rho * cp
    theta = 0.5
    steps = int(round(t_end / dt))
    per_second = int(round(1.0 / dt))
    if abs(per_second * dt - 1.0) > 1e-10:
        raise ValueError("dt must divide one second")

    temp = np.full(n, 28.0)
    moist = np.full(n, 2.55)
    # outputs include interior nodes and reconstructed surface at requested grid
    tout = np.empty((int(round(t_end)), n + 1))
    cout = np.empty_like(tout)
    I = eye(n, format="csc")
    coeff_t = np.full(n + 1, k)
    At, _, _ = assemble(r, coeff_t, h, 0.0, cap)
    Lt = (I - theta * dt * At.tocsc()).tocsc()
    Rt = (I + (1.0 - theta) * dt * At.tocsc()).tocsc()

    # Initial moisture surface and operator.
    c_surf = surface_from_inner(moist, 7e-9 * np.exp(-0.89 / 2.55), hm, ac[0], dr)
    c_full = np.r_[moist, c_surf]
    Ac_old, bc_old, _ = assemble(r, 7e-9 * np.exp(-0.89 / np.maximum(c_full, 1e-12)), hm, ac[0], 1.0)
    out = 0
    for step in range(steps):
        t0, t1 = step * dt, (step + 1) * dt
        ta0, ta1 = float(np.interp(t0, at, aa)), float(np.interp(t1, at, aa))
        ca0, ca1 = float(np.interp(t0, at, ac)), float(np.interp(t1, at, ac))
        _, bt0, _ = assemble(r, coeff_t, h, ta0, cap)
        _, bt1, _ = assemble(r, coeff_t, h, ta1, cap)
        temp = np.asarray(spsolve(Lt, Rt @ temp + dt * (theta * bt1 + (1.0 - theta) * bt0)))

        rhs = (I + (1.0 - theta) * dt * Ac_old.tocsc()) @ moist + dt * (1.0 - theta) * bc_old
        qi = moist.copy()
        cs_i = c_surf
        for _ in range(40):
            d_full = 7e-9 * np.exp(-0.89 / np.maximum(np.r_[qi, cs_i], 1e-12))
            An, bn, _ = assemble(r, d_full, hm, ca1, 1.0)
            Ln = (I - theta * dt * An.tocsc()).tocsc()
            cand = np.asarray(spsolve(Ln, rhs + dt * theta * bn))
            # Update surface coefficient and reconstructed value consistently.
            ds = 7e-9 * np.exp(-0.89 / max(cs_i, 1e-12))
            cs_new = surface_from_inner(cand, ds, hm, ca1, dr)
            qi_new = 0.7 * cand + 0.3 * qi
            cs_new = 0.7 * cs_new + 0.3 * cs_i
            if max(np.max(np.abs(qi_new - qi)), abs(cs_new - cs_i)) < 1e-12:
                qi, cs_i = qi_new, cs_new
                break
            qi, cs_i = qi_new, cs_new
        moist, c_surf = qi, cs_i
        c_full = np.r_[moist, c_surf]
        d_full = 7e-9 * np.exp(-0.89 / np.maximum(c_full, 1e-12))
        Ac_old, bc_old, _ = assemble(r, d_full, hm, ca1, 1.0)

        if (step + 1) % per_second == 0:
            tout[out, :-1] = temp
            tout[out, -1] = surface_from_inner(temp, k, h, ta1, dr)
            cout[out, :-1] = moist
            cout[out, -1] = c_surf
            out += 1
    return np.arange(1, int(round(t_end)) + 1), r, tout, cout


def extract(arr, r, t):
    """Extract values at 0,0.5,1,1.5,2 cm; arr includes surface."""
    targets = np.array([0, 0.5, 1, 1.5, 2.0]) / 100.0
    vals = []
    for x in targets:
        if x == 0:
            vals.append((9 * arr[t - 1, 0] - arr[t - 1, 1]) / 8.0)
        elif abs(x - r[-1]) < 1e-14:
            vals.append(arr[t - 1, -1])
        else:
            vals.append(np.interp(x, r, arr[t - 1]))
    return np.asarray(vals)


if __name__ == "__main__":
    times, r, T, C = solve_q1(dr_cm=0.025, dt=0.25)
    req = [100, 300, 600, 900, 1200, 1500, 1800]
    print("TABLE1")
    for t in req: print(t, " ".join(f"{v:.4f}" for v in extract(T, r, t)))
    print("TABLE2")
    for t in req: print(t, " ".join(f"{v:.4f}" for v in extract(C, r, t)))
