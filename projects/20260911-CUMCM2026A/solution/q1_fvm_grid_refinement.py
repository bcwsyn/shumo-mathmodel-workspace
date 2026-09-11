"""Problem 1: conservative radial FVM with grid-refinement verification.

Outputs
-------
* result1_FVM_refined.xlsx: final 1 s x 0.1 cm temperature/moisture fields.
* q1_grid_convergence.xlsx: A/B/C/D grid comparison at the 35 required points.
* q1_grid_convergence.json: machine-readable run and physical-check summary.

The spatial convergence runs use the same dt=0.25 s, so only the radial mesh
changes.  The delivered field additionally uses dt=0.125 s on the finest mesh.
"""

from __future__ import annotations

import json
import math
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from scipy.linalg import solve_banded


ROOT = Path(__file__).resolve().parents[1]
AIR_FILE = ROOT / "附件" / "附件1.xlsx"
TEMPLATE_FILE = ROOT / "附件" / "附件3" / "result1.xlsx"
FINAL_FILE = ROOT / "附件" / "附件3" / "result1_FVM_refined.xlsx"
CONVERGENCE_FILE = ROOT / "solution" / "q1_grid_convergence.xlsx"
SUMMARY_FILE = ROOT / "solution" / "q1_grid_convergence.json"

RADIUS_M = 0.02
RHO = 820.0
CP = 2600.0
K = 0.36
H = 25.0
HM = 8e-7
THETA = 0.5  # Crank--Nicolson

KEY_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=int)
KEY_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
TYPICAL_TIMES = (100, 600, 1800)
TYPICAL_RADII_CM = (0.0, 1.0, 2.0)


@dataclass
class GridResult:
    label: str
    dr_cm: float
    dt_s: float
    times_s: np.ndarray
    radii_cm: np.ndarray
    temperature: np.ndarray
    moisture: np.ndarray
    runtime_s: float
    checks: dict[str, float | int | bool | str]


def read_air_data() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(AIR_FILE, read_only=True, data_only=True)
    data = np.asarray(list(wb.active.iter_rows(min_row=2, values_only=True)), dtype=float)
    wb.close()
    if data.shape != (241, 3):
        raise ValueError(f"附件1尺寸异常：期望(241, 3)，实际{data.shape}")
    if not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError("附件1时间列必须严格递增")
    return data[:, 0], data[:, 1], data[:, 2]


def moisture_diffusivity(c: np.ndarray | float) -> np.ndarray | float:
    return 7e-9 * np.exp(-0.89 / np.maximum(c, 1e-12))


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    return 2.0 * left * right / (left + right)


def reconstruct_surface(
    last_cell_value: float,
    surface_coefficient: float,
    transfer_coefficient: float,
    air_value: float,
    half_cell_width: float,
) -> float:
    """Reconstruct the Robin-boundary surface value from the last cell."""
    return (
        surface_coefficient * last_cell_value
        + half_cell_width * transfer_coefficient * air_value
    ) / (surface_coefficient + half_cell_width * transfer_coefficient)


def build_operator(
    cell_centers: np.ndarray,
    cell_volumes: np.ndarray,
    face_radii: np.ndarray,
    cell_coefficients: np.ndarray,
    surface_coefficient: float,
    transfer_coefficient: float,
    air_value: float,
    capacity: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, float]:
    """Return tridiagonal A and b for dq/dt=Aq+b.

    Integrating (1/r)d/dr(r Gamma dq/dr) over each annular control volume
    makes all internal fluxes cancel pairwise.  The center face has zero area,
    and the outer face uses a Robin resistance over half a cell.
    """
    m = len(cell_centers)
    dr = cell_centers[1] - cell_centers[0]
    internal_conductance = (
        face_radii[1:-1]
        * harmonic_mean(cell_coefficients[:-1], cell_coefficients[1:])
        / dr
    )

    lower = np.zeros(m, dtype=float)
    diagonal = np.zeros(m, dtype=float)
    upper = np.zeros(m, dtype=float)
    source = np.zeros(m, dtype=float)

    lower[1:] = internal_conductance / (cell_volumes[1:] * capacity)
    upper[:-1] = internal_conductance / (cell_volumes[:-1] * capacity)
    diagonal -= lower + upper

    half = dr / 2.0
    boundary_conductance = (
        RADIUS_M
        * transfer_coefficient
        * surface_coefficient
        / (surface_coefficient + half * transfer_coefficient)
    )
    diagonal[-1] -= boundary_conductance / (cell_volumes[-1] * capacity)
    source[-1] = boundary_conductance * air_value / (cell_volumes[-1] * capacity)
    return lower, diagonal, upper, source, boundary_conductance


def tridiagonal_matvec(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray, vector: np.ndarray
) -> np.ndarray:
    result = diagonal * vector
    result[1:] += lower[1:] * vector[:-1]
    result[:-1] += upper[:-1] * vector[1:]
    return result


def solve_implicit_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    rhs: np.ndarray,
    factor: float,
) -> np.ndarray:
    """Solve (I-factor*A)x=rhs for a tridiagonal A."""
    band = np.zeros((3, len(diagonal)), dtype=float)
    band[0, 1:] = -factor * upper[:-1]
    band[1] = 1.0 - factor * diagonal
    band[2, :-1] = -factor * lower[1:]
    return solve_banded((1, 1), band, rhs, check_finite=False)


def reconstruct_profile(
    cell_values: np.ndarray,
    cell_centers: np.ndarray,
    surface_value: float,
    output_radii_m: np.ndarray,
) -> np.ndarray:
    """Interpolate cell-centred values to r=0,0.1,...,2.0 cm."""
    center_value = (9.0 * cell_values[0] - cell_values[1]) / 8.0
    locations = np.r_[0.0, cell_centers, RADIUS_M]
    values = np.r_[center_value, cell_values, surface_value]
    return np.interp(output_radii_m, locations, values)


def solve_grid(label: str, dr_cm: float, dt_s: float) -> GridResult:
    started = time.perf_counter()
    air_time, air_temperature, air_moisture = read_air_data()
    cells = int(round(2.0 / dr_cm))
    dr = RADIUS_M / cells
    if not math.isclose(dr_cm, dr * 100.0, rel_tol=0.0, abs_tol=1e-12):
        raise ValueError("空间步长必须整除2 cm半径")
    steps_per_second = int(round(1.0 / dt_s))
    if not math.isclose(steps_per_second * dt_s, 1.0, abs_tol=1e-12):
        raise ValueError("时间步长必须整除1 s")

    faces = np.arange(cells + 1, dtype=float) * dr
    centers = (np.arange(cells, dtype=float) + 0.5) * dr
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    output_radii_cm = np.arange(0.0, 2.0000001, 0.1)
    output_radii_m = output_radii_cm / 100.0
    times_s = np.arange(1, 1801, dtype=int)
    temperature_out = np.empty((1800, 21), dtype=float)
    moisture_out = np.empty_like(temperature_out)

    temperature = np.full(cells, 28.0, dtype=float)
    moisture = np.full(cells, 2.55, dtype=float)
    half = dr / 2.0

    # At t=0 the temperature boundary is compatible with the initial field.
    surface_temperature = 28.0
    surface_moisture = 2.55
    for _ in range(100):
        ds = float(moisture_diffusivity(surface_moisture))
        updated = reconstruct_surface(moisture[-1], ds, HM, air_moisture[0], half)
        if abs(updated - surface_moisture) < 1e-14:
            surface_moisture = updated
            break
        surface_moisture = 0.7 * updated + 0.3 * surface_moisture

    constant_k = np.full(cells, K, dtype=float)
    lt, dt_diag, ut, _, beta_t = build_operator(
        centers, volumes, faces, constant_k, K, H, 0.0, RHO * CP
    )

    dcells = np.asarray(moisture_diffusivity(moisture), dtype=float)
    dsurface = float(moisture_diffusivity(surface_moisture))
    lc_old, dc_old, uc_old, bc_old, beta_c_old = build_operator(
        centers, volumes, faces, dcells, dsurface, HM, air_moisture[0], 1.0
    )

    heat_balance_error = 0.0
    moisture_balance_error = 0.0
    heat_flux_integral = 0.0
    moisture_flux_integral = 0.0
    initial_heat_storage = float(np.sum(volumes * RHO * CP * temperature))
    initial_moisture_storage = float(np.sum(volumes * moisture))
    max_picard_iterations = 0
    total_steps = int(round(1800.0 / dt_s))
    output_index = 0

    for step in range(total_steps):
        t0 = step * dt_s
        t1 = (step + 1) * dt_s
        ta0 = float(np.interp(t0, air_time, air_temperature))
        ta1 = float(np.interp(t1, air_time, air_temperature))
        ca0 = float(np.interp(t0, air_time, air_moisture))
        ca1 = float(np.interp(t1, air_time, air_moisture))

        # Temperature: constant operator, time-varying boundary source only.
        source_t0 = np.zeros(cells); source_t0[-1] = beta_t * ta0 / (volumes[-1] * RHO * CP)
        source_t1 = np.zeros(cells); source_t1[-1] = beta_t * ta1 / (volumes[-1] * RHO * CP)
        rhs_t = (
            temperature
            + (1.0 - THETA) * dt_s * tridiagonal_matvec(lt, dt_diag, ut, temperature)
            + dt_s * ((1.0 - THETA) * source_t0 + THETA * source_t1)
        )
        old_temperature = temperature
        old_surface_temperature = surface_temperature
        temperature = solve_implicit_tridiagonal(
            lt, dt_diag, ut, rhs_t, THETA * dt_s
        )
        surface_temperature = reconstruct_surface(temperature[-1], K, H, ta1, half)

        # Moisture: Picard iteration for D(C), including D at the surface.
        rhs_c = (
            moisture
            + (1.0 - THETA) * dt_s
            * tridiagonal_matvec(lc_old, dc_old, uc_old, moisture)
            + (1.0 - THETA) * dt_s * bc_old
        )
        old_moisture = moisture
        old_surface_moisture = surface_moisture
        iterate = moisture.copy()
        iterate_surface = surface_moisture
        converged = False
        for picard_iteration in range(1, 51):
            dcells = np.asarray(moisture_diffusivity(iterate), dtype=float)
            dsurface = float(moisture_diffusivity(iterate_surface))
            lc, dc, uc, bc, beta_c = build_operator(
                centers, volumes, faces, dcells, dsurface, HM, ca1, 1.0
            )
            candidate = solve_implicit_tridiagonal(
                lc, dc, uc, rhs_c + THETA * dt_s * bc, THETA * dt_s
            )
            candidate_surface = reconstruct_surface(
                candidate[-1], dsurface, HM, ca1, half
            )
            updated = 0.8 * candidate + 0.2 * iterate
            updated_surface = 0.8 * candidate_surface + 0.2 * iterate_surface
            error = max(
                float(np.max(np.abs(updated - iterate))),
                abs(updated_surface - iterate_surface),
            )
            iterate = updated
            iterate_surface = updated_surface
            if error < 1e-11:
                converged = True
                break
        if not converged:
            raise RuntimeError(f"{label}网格在t={t1:.3f}s的Picard迭代未收敛")
        max_picard_iterations = max(max_picard_iterations, picard_iteration)
        moisture = iterate
        surface_moisture = iterate_surface

        # Rebuild the converged operator for the next CN old-time term.
        dcells = np.asarray(moisture_diffusivity(moisture), dtype=float)
        dsurface = float(moisture_diffusivity(surface_moisture))
        lc_old, dc_old, uc_old, bc_old, beta_c_old = build_operator(
            centers, volumes, faces, dcells, dsurface, HM, ca1, 1.0
        )

        # Global FVM balance. Internal face fluxes cancel exactly.
        heat_flux0 = RADIUS_M * H * (ta0 - old_surface_temperature)
        heat_flux1 = RADIUS_M * H * (ta1 - surface_temperature)
        moisture_flux0 = RADIUS_M * HM * (ca0 - old_surface_moisture)
        moisture_flux1 = RADIUS_M * HM * (ca1 - surface_moisture)
        heat_flux_integral += dt_s * ((1.0 - THETA) * heat_flux0 + THETA * heat_flux1)
        moisture_flux_integral += dt_s * (
            (1.0 - THETA) * moisture_flux0 + THETA * moisture_flux1
        )
        heat_storage = float(np.sum(volumes * RHO * CP * temperature))
        moisture_storage = float(np.sum(volumes * moisture))
        heat_balance_error = max(
            heat_balance_error,
            abs((heat_storage - initial_heat_storage) - heat_flux_integral),
        )
        moisture_balance_error = max(
            moisture_balance_error,
            abs((moisture_storage - initial_moisture_storage) - moisture_flux_integral),
        )

        if (step + 1) % steps_per_second == 0:
            temperature_out[output_index] = reconstruct_profile(
                temperature, centers, surface_temperature, output_radii_m
            )
            moisture_out[output_index] = reconstruct_profile(
                moisture, centers, surface_moisture, output_radii_m
            )
            output_index += 1

    air_t_at_seconds = np.interp(times_s, air_time, air_temperature)
    air_c_at_seconds = np.interp(times_s, air_time, air_moisture)
    checks: dict[str, float | int | bool | str] = {
        "cells": cells,
        "temperature_min": float(temperature_out.min()),
        "temperature_max": float(temperature_out.max()),
        "moisture_min": float(moisture_out.min()),
        "moisture_max": float(moisture_out.max()),
        "temperature_radial_violation_count": int(np.count_nonzero(np.diff(temperature_out, axis=1) < -1e-8)),
        "moisture_radial_violation_count": int(np.count_nonzero(np.diff(moisture_out, axis=1) > 1e-8)),
        "negative_moisture_count": int(np.count_nonzero(moisture_out < -1e-12)),
        "surface_hotter_than_air_count": int(np.count_nonzero(temperature_out[:, -1] > air_t_at_seconds + 1e-8)),
        "surface_drier_than_air_count": int(np.count_nonzero(moisture_out[:, -1] < air_c_at_seconds - 1e-8)),
        "max_picard_iterations": max_picard_iterations,
        "heat_global_balance_abs_error": heat_balance_error,
        "moisture_global_balance_abs_error": moisture_balance_error,
        "physical_checks_passed": False,
    }
    checks["physical_checks_passed"] = all(
        checks[key] == 0
        for key in (
            "temperature_radial_violation_count",
            "moisture_radial_violation_count",
            "negative_moisture_count",
            "surface_hotter_than_air_count",
            "surface_drier_than_air_count",
        )
    )
    return GridResult(
        label=label,
        dr_cm=dr_cm,
        dt_s=dt_s,
        times_s=times_s,
        radii_cm=output_radii_cm,
        temperature=temperature_out,
        moisture=moisture_out,
        runtime_s=time.perf_counter() - started,
        checks=checks,
    )


def key_values(result: GridResult, field: str) -> np.ndarray:
    data = result.temperature if field == "temperature" else result.moisture
    time_indices = KEY_TIMES - 1
    radius_indices = np.rint(KEY_RADII_CM / 0.1).astype(int)
    return data[np.ix_(time_indices, radius_indices)]


def error_summary(coarse: GridResult, fine: GridResult, field: str) -> dict[str, float | int]:
    a = key_values(coarse, field)
    b = key_values(fine, field)
    difference = np.abs(a - b)
    relative = difference / np.maximum(np.abs(b), 1e-15) * 100.0
    worst = np.unravel_index(int(np.argmax(difference)), difference.shape)
    return {
        "max_abs": float(difference[worst]),
        "mean_abs": float(difference.mean()),
        "max_rel_percent": float(relative.max()),
        "worst_time_s": int(KEY_TIMES[worst[0]]),
        "worst_radius_cm": float(KEY_RADII_CM[worst[1]]),
    }


def save_result1(final: GridResult) -> None:
    wb = load_workbook(TEMPLATE_FILE)
    if len(wb.worksheets) < 2:
        raise ValueError("result1模板缺少温度或水分浓度工作表")
    for ws, data in zip(wb.worksheets[:2], (final.temperature, final.moisture)):
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        if ws.max_column > 1:
            ws.delete_cols(2, ws.max_column - 1)
        ws.cell(1, 1).value = "时间\\到药材中心的距离"
        for column, radius in enumerate(final.radii_cm, start=2):
            ws.cell(1, column).value = round(float(radius), 1)
        for row, second in enumerate(final.times_s, start=2):
            ws.cell(row, 1).value = int(second)
            for column, value in enumerate(data[row - 2], start=2):
                cell = ws.cell(row, column)
                cell.value = round(float(value), 4)
                cell.number_format = "0.0000"
        ws.freeze_panes = "B2"
    wb.save(FINAL_FILE)


def append_comparison_sheet(
    wb: Workbook,
    title: str,
    field: str,
    results: list[GridResult],
    typical_only: bool,
) -> None:
    ws = wb.create_sheet(title)
    headings = [
        "时间/s", "距离/cm",
        "Δr=0.1 cm", "Δr=0.05 cm", "Δr=0.025 cm", "Δr=0.0125 cm",
        "|A-B|", "|B-C|", "|C-D|", "C→D相对误差/%",
    ]
    ws.append(headings)
    times = TYPICAL_TIMES if typical_only else KEY_TIMES
    radii = TYPICAL_RADII_CM if typical_only else KEY_RADII_CM
    arrays = [r.temperature if field == "temperature" else r.moisture for r in results]
    for second in times:
        for radius in radii:
            ti = int(second) - 1
            ri = int(round(float(radius) / 0.1))
            values = [float(a[ti, ri]) for a in arrays]
            ws.append([
                int(second), float(radius), *values,
                abs(values[0] - values[1]),
                abs(values[1] - values[2]),
                abs(values[2] - values[3]),
                abs(values[2] - values[3]) / max(abs(values[3]), 1e-15) * 100.0,
            ])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=9):
        for cell in row:
            cell.number_format = "0.000000"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def save_convergence_workbook(
    spatial_results: list[GridResult],
    final_result: GridResult,
    summaries: dict[str, dict[str, dict[str, float | int]]],
    temporal: dict[str, dict[str, float | int]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["变量", "比较", "最大绝对误差", "平均绝对误差", "最大相对误差/%", "最差时间/s", "最差距离/cm", "判定"])
    for field, chinese in (("temperature", "温度"), ("moisture", "水分浓度")):
        for comparison in ("A-B", "B-C", "C-D"):
            item = summaries[field][comparison]
            passed = item["max_abs"] < 1e-4 or item["max_rel_percent"] < 0.1
            ws.append([
                chinese, comparison, item["max_abs"], item["mean_abs"],
                item["max_rel_percent"], item["worst_time_s"],
                item["worst_radius_cm"], "通过" if passed else "继续加密",
            ])
        item = temporal[field]
        ws.append([
            chinese, "D网格 Δt=0.25→0.125 s", item["max_abs"], item["mean_abs"],
            item["max_rel_percent"], item["worst_time_s"], item["worst_radius_cm"],
            "时间离散稳定" if item["max_rel_percent"] < 0.1 else "需继续减小时间步长",
        ])
    ws.append([])
    ws.append(["最终采用", "Δr=0.0125 cm，Δt=0.125 s；输出重采样至0.1 cm和1 s"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "0.000000"
    ws.freeze_panes = "A2"

    append_comparison_sheet(wb, "温度_典型9点", "temperature", spatial_results, True)
    append_comparison_sheet(wb, "水分_典型9点", "moisture", spatial_results, True)
    append_comparison_sheet(wb, "温度_完整35点", "temperature", spatial_results, False)
    append_comparison_sheet(wb, "水分_完整35点", "moisture", spatial_results, False)

    check_ws = wb.create_sheet("物理与守恒检查")
    keys = list(final_result.checks.keys())
    check_ws.append(["网格", "Δr/cm", "Δt/s", "运行时间/s", *keys])
    for result in spatial_results + [final_result]:
        check_ws.append([
            result.label, result.dr_cm, result.dt_s, result.runtime_s,
            *[result.checks[k] for k in keys],
        ])
    for cell in check_ws[1]:
        cell.font = Font(bold=True)
    check_ws.freeze_panes = "A2"
    wb.save(CONVERGENCE_FILE)


def main() -> None:
    spatial_specs = [
        ("A", 0.1, 0.25),
        ("B", 0.05, 0.25),
        ("C", 0.025, 0.25),
        ("D", 0.0125, 0.25),
    ]
    spatial_results = []
    for label, dr_cm, dt_s in spatial_specs:
        print(f"RUN grid={label} dr={dr_cm}cm dt={dt_s}s", flush=True)
        result = solve_grid(label, dr_cm, dt_s)
        spatial_results.append(result)
        print(
            f"DONE grid={label} runtime={result.runtime_s:.2f}s "
            f"physical_checks={result.checks['physical_checks_passed']}",
            flush=True,
        )

    print("RUN final grid=D dr=0.0125cm dt=0.125s", flush=True)
    final_result = solve_grid("D-final", 0.0125, 0.125)
    print(f"DONE final runtime={final_result.runtime_s:.2f}s", flush=True)

    comparisons = (("A-B", 0, 1), ("B-C", 1, 2), ("C-D", 2, 3))
    summaries: dict[str, dict[str, dict[str, float | int]]] = {}
    for field in ("temperature", "moisture"):
        summaries[field] = {
            name: error_summary(spatial_results[i], spatial_results[j], field)
            for name, i, j in comparisons
        }
    temporal = {
        field: error_summary(spatial_results[-1], final_result, field)
        for field in ("temperature", "moisture")
    }

    save_result1(final_result)
    save_convergence_workbook(spatial_results, final_result, summaries, temporal)

    report = {
        "method": "conservative cell-centred radial finite-volume method",
        "time_method": "Crank-Nicolson with Picard iteration for D(C)",
        "spatial_runs": [
            {
                "label": r.label, "dr_cm": r.dr_cm, "dt_s": r.dt_s,
                "runtime_s": r.runtime_s, "checks": r.checks,
            }
            for r in spatial_results
        ],
        "spatial_error_at_35_key_points": summaries,
        "temporal_error_at_35_key_points": temporal,
        "final_run": {
            "dr_cm": final_result.dr_cm,
            "dt_s": final_result.dt_s,
            "runtime_s": final_result.runtime_s,
            "checks": final_result.checks,
        },
        "outputs": {
            "result1": str(FINAL_FILE),
            "convergence_workbook": str(CONVERGENCE_FILE),
        },
    }
    SUMMARY_FILE.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print("GRID_SUMMARY")
    print(json.dumps(summaries, ensure_ascii=False))
    print("TEMPORAL_SUMMARY")
    print(json.dumps(temporal, ensure_ascii=False))
    print(f"RESULT1={FINAL_FILE}")
    print(f"CONVERGENCE={CONVERGENCE_FILE}")


if __name__ == "__main__":
    main()

