"""Run the cell-centred Q1 solver and fill result1.xlsx."""
from __future__ import annotations
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from test_cell_q1 import solve, surface_value, read_air

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "附件" / "附件3" / "result1.xlsx"
OUT = ROOT / "附件" / "附件3" / "result1_computed.xlsx"
SUMMARY = Path(__file__).resolve().parent / "q1_cell_run_summary.json"


def samples(arr, centers, t_index, surface, target_cm):
    x = np.asarray(target_cm, float) / 100.0
    row = arr[t_index]
    out = []
    dr = centers[1] - centers[0]
    for xx in x:
        if abs(xx) < 1e-15:
            out.append((9.0 * row[0] - row[1]) / 8.0)
        elif abs(xx - (centers[-1] + dr / 2.0)) < 1e-12:
            out.append(surface)
        else:
            out.append(float(np.interp(xx, centers, row)))
    return np.asarray(out)


def build_outputs(dr_cm=0.025, dt=0.25):
    times, centers, Tcell, Ccell = solve(dr_cm, dt)
    at, aa, ac = read_air()
    targets = np.arange(0.0, 2.00001, 0.1)
    T = np.empty((len(times), len(targets)))
    C = np.empty_like(T)
    dr = centers[1] - centers[0]
    for i, tv in enumerate(times.astype(int)):
        ta = float(np.interp(tv, at, aa)); ca = float(np.interp(tv, at, ac))
        ts = surface_value(Tcell[i, -1], 0.36, 25.0, ta, dr)
        ds = 7e-9 * np.exp(-0.89 / max(Ccell[i, -1], 1e-12))
        cs = surface_value(Ccell[i, -1], ds, 8e-7, ca, dr)
        T[i] = samples(Tcell, centers, i, ts, targets)
        C[i] = samples(Ccell, centers, i, cs, targets)
    return times, targets, T, C


def write_xlsx(times, targets, T, C):
    wb = load_workbook(TEMPLATE)
    actual = OUT
    for ws, data in zip(wb.worksheets[:2], (T, C)):
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ws.cell(1, 1).value = "时间\\到药材中心的距离"
        for j, x in enumerate(targets, start=2):
            ws.cell(1, j).value = round(float(x), 10)
        for i, tv in enumerate(times, start=2):
            ws.cell(i, 1).value = int(tv)
            for j, v in enumerate(data[i - 2], start=2):
                ws.cell(i, j).value = float(v)
                ws.cell(i, j).number_format = "0.0000"
        ws.freeze_panes = "B2"
    try:
        wb.save(TEMPLATE)
        actual = TEMPLATE
    except PermissionError:
        wb.save(OUT)
    return actual


def print_table(name, data, times):
    req_t = [100, 300, 600, 900, 1200, 1500, 1800]
    req_r = [0, 0.5, 1.0, 1.5, 2.0]
    idx = {int(v): i for i, v in enumerate(times)}
    ridx = [int(round(x / 0.1)) for x in req_r]
    print(name)
    print("time_s\t" + "\t".join(f"r={x:g}cm" for x in req_r))
    for tv in req_t:
        print(tv, "\t" + "\t".join(f"{data[idx[tv],j]:.4f}" for j in ridx))


def main():
    # Use a finer internal mesh for the delivered four-decimal table.
    times, targets, T, C = build_outputs(0.0125, 0.125)
    # Compare at the actual requested physical output points, not raw cell indices.
    tc, xc, Tc_cell, Cc_cell = build_outputs(0.025, 0.25)
    diffT = float(np.max(np.abs(T[[99,299,599,899,1199,1499,1799]] - Tc_cell[[99,299,599,899,1199,1499,1799]])))
    diffC = float(np.max(np.abs(C[[99,299,599,899,1199,1499,1799]] - Cc_cell[[99,299,599,899,1199,1499,1799]])))
    actual = write_xlsx(times, targets, T, C)
    summary = {
        "solver": "cell-centred radial finite-volume Crank-Nicolson with implicit Picard moisture update",
        "dr_internal_cm": 0.0125, "dt_internal_s": 0.125,
        "output_shape": [int(T.shape[0]), int(T.shape[1])],
        "max_fine_vs_coarse_T_at_requested_points": diffT,
        "max_fine_vs_coarse_C_at_requested_points": diffC,
        "min_T": float(T.min()), "max_T": float(T.max()),
        "min_C": float(C.min()), "max_C": float(C.max()),
        "result_file": str(actual),
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print_table("表1 温度 (°C)", T, times)
    print_table("表2 水分浓度 (kg/kg)", C, times)
    print("RUN_SUMMARY")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
