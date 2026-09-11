"""Problem 1 solver for CUMCM 2026 A.

Axisymmetric radial heat conduction and nonlinear moisture diffusion in a
fixed cylinder.  The script fills the supplied result1.xlsx template and
prints the requested Table 1/Table 2 entries.
"""

from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.sparse import eye, lil_matrix
from scipy.sparse.linalg import spsolve


ROOT = Path(__file__).resolve().parents[1]
AIR_FILE = ROOT / "附件" / "附件1.xlsx"
TEMPLATE_FILE = ROOT / "附件" / "附件3" / "result1.xlsx"
OUTPUT_FILE = ROOT / "附件" / "附件3" / "result1.xlsx"
FALLBACK_OUTPUT_FILE = ROOT / "附件" / "附件3" / "result1_computed.xlsx"
SUMMARY_FILE = Path(__file__).resolve().parent / "q1_run_summary.json"


def read_air() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(AIR_FILE, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    wb.close()
    arr = np.asarray(rows, dtype=float)
    return arr[:, 0], arr[:, 1], arr[:, 2]


def harmonic(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    den = a + b
    return np.divide(2.0 * a * b, den, out=np.zeros_like(den), where=den > 0)


def assemble_operator(
    r: np.ndarray,
    coeff: np.ndarray,
    robin_h: float,
    boundary_value: float,
    capacity: float,
) -> tuple[lil_matrix, np.ndarray]:
    """Build q_t = A q + b using radial finite volumes.

    Nodes include r=0 and r=R.  The first and last control volumes are
    half-width; the center limit gives the correct cylindrical factor 4.
    """
    n = len(r) - 1
    dr = r[1] - r[0]
    radius = r[-1]
    faces = r[:-1] + 0.5 * dr
    face_coeff = harmonic(coeff[:-1], coeff[1:])
    A = lil_matrix((n + 1, n + 1), dtype=float)
    b = np.zeros(n + 1, dtype=float)

    # Center control volume: V = dr^2/8 and zero flux at r=0.
    v0 = dr * dr / 8.0
    c0 = faces[0] * face_coeff[0] / dr
    A[0, 0] = -c0 / (v0 * capacity)
    A[0, 1] = c0 / (v0 * capacity)

    # Interior control volumes.
    for i in range(1, n):
        vol = r[i] * dr
        cm = faces[i - 1] * face_coeff[i - 1] / dr
        cp = faces[i] * face_coeff[i] / dr
        A[i, i - 1] = cm / (vol * capacity)
        A[i, i] = -(cm + cp) / (vol * capacity)
        A[i, i + 1] = cp / (vol * capacity)

    # Surface control volume.  F_R = R*h*(air - surface).
    vol = radius * dr / 2.0 - dr * dr / 8.0
    cm = faces[-1] * face_coeff[-1] / dr
    A[n, n - 1] = cm / (vol * capacity)
    A[n, n] = -(cm + radius * robin_h) / (vol * capacity)
    b[n] = radius * robin_h * boundary_value / (vol * capacity)
    return A, b


def solve_q1(dr_cm: float = 0.05, dt: float = 0.5, t_end: float = 1800.0):
    """Solve with Crank--Nicolson and return integer-second snapshots."""
    air_t, air_temp, air_moist = read_air()
    dr = dr_cm / 100.0
    n = int(round(0.02 / dr))
    r = np.arange(n + 1, dtype=float) * dr
    if abs(r[-1] - 0.02) > 1e-12:
        raise ValueError("dr must divide the 2 cm radius exactly")

    rho, cp, k, h, hm = 820.0, 2600.0, 0.36, 25.0, 8e-7
    cap = rho * cp
    theta = 0.5
    n_steps = int(round(t_end / dt))
    substeps_per_second = int(round(1.0 / dt))
    if abs(substeps_per_second * dt - 1.0) > 1e-10:
        raise ValueError("dt must divide one second for the requested output")

    temp = np.full(n + 1, 28.0, dtype=float)
    moist = np.full(n + 1, 2.55, dtype=float)
    times_out = np.arange(1, int(round(t_end)) + 1, dtype=float)
    temp_out = np.empty((len(times_out), n + 1), dtype=float)
    moist_out = np.empty_like(temp_out)

    coeff_t = np.full(n + 1, k, dtype=float)
    A_t, _ = assemble_operator(r, coeff_t, h, 0.0, cap)
    I = eye(n + 1, format="csc")
    lhs_t = (I - theta * dt * A_t.tocsc()).tocsc()
    rhs_t_mat = (I + (1.0 - theta) * dt * A_t.tocsc()).tocsc()

    # Initial moisture operator for the first Crank--Nicolson right-hand side.
    D0 = 7e-9 * np.exp(-0.89 / moist)
    A_c_old, b_c_old = assemble_operator(r, D0, hm, float(air_moist[0]), 1.0)

    out_idx = 0
    for step in range(n_steps):
        t0 = step * dt
        t1 = (step + 1) * dt
        ta0 = float(np.interp(t0, air_t, air_temp))
        ta1 = float(np.interp(t1, air_t, air_temp))
        ca0 = float(np.interp(t0, air_t, air_moist))
        ca1 = float(np.interp(t1, air_t, air_moist))

        # Temperature is linear for Problem 1 and can be solved directly.
        _, b_t0 = assemble_operator(r, coeff_t, h, ta0, cap)
        _, b_t1 = assemble_operator(r, coeff_t, h, ta1, cap)
        rhs_t = rhs_t_mat @ temp + dt * (theta * b_t1 + (1.0 - theta) * b_t0)
        temp = np.asarray(spsolve(lhs_t, rhs_t), dtype=float)

        # Moisture has D(C); Picard iteration handles the implicit CN step.
        rhs_c = (
            (I + (1.0 - theta) * dt * A_c_old.tocsc()) @ moist
            + dt * (1.0 - theta) * b_c_old
        )
        moist_iter = moist.copy()
        for _ in range(30):
            D_iter = 7e-9 * np.exp(-0.89 / np.maximum(moist_iter, 1e-12))
            A_c_new, b_c_new = assemble_operator(r, D_iter, hm, ca1, 1.0)
            lhs_c = (I - theta * dt * A_c_new.tocsc()).tocsc()
            candidate = np.asarray(spsolve(lhs_c, rhs_c + dt * theta * b_c_new), dtype=float)
            # Mild relaxation prevents a noisy boundary input from hurting
            # convergence while preserving the implicit fixed point.
            updated = 0.7 * candidate + 0.3 * moist_iter
            if np.max(np.abs(updated - moist_iter)) < 1e-12:
                moist_iter = updated
                break
            moist_iter = updated
        moist = moist_iter
        A_c_old, b_c_old = assemble_operator(
            r,
            7e-9 * np.exp(-0.89 / np.maximum(moist, 1e-12)),
            hm,
            ca1,
            1.0,
        )

        if (step + 1) % substeps_per_second == 0:
            temp_out[out_idx] = temp
            moist_out[out_idx] = moist
            out_idx += 1

    return times_out, r, temp_out, moist_out


def write_result1(times: np.ndarray, r: np.ndarray, temp: np.ndarray, moist: np.ndarray) -> Path:
    # The fine internal grid is used for accuracy; the supplied template
    # requires output at 0.1 cm, so retain every second node here.
    output_stride = int(round((0.1 / 100.0) / (r[1] - r[0])))
    output_idx = np.arange(0, len(r), output_stride, dtype=int)
    r_write = r[output_idx]
    temp_write = temp[:, output_idx]
    moist_write = moist[:, output_idx]
    wb = load_workbook(TEMPLATE_FILE)
    sheets = wb.worksheets
    if len(sheets) < 2:
        raise ValueError("result1.xlsx must contain temperature and moisture sheets")
    for ws, data in ((sheets[0], temp_write), (sheets[1], moist_write)):
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ws.cell(1, 1).value = "时间\\到药材中心的距离"
        for j, rv in enumerate(r_write * 100.0, start=2):
            ws.cell(1, j).value = round(float(rv), 10)
        for i, (tv, row_values) in enumerate(zip(times, data), start=2):
            ws.cell(i, 1).value = int(round(float(tv)))
            for j, value in enumerate(row_values, start=2):
                cell = ws.cell(i, j)
                cell.value = float(value)
                cell.number_format = "0.0000"
        ws.freeze_panes = "B2"
    try:
        wb.save(OUTPUT_FILE)
        return OUTPUT_FILE
    except PermissionError:
        # The supplied template may be open in WPS/Excel.  Preserve the
        # computed result without closing or killing the user's application.
        wb.save(FALLBACK_OUTPUT_FILE)
        return FALLBACK_OUTPUT_FILE


def format_table(title: str, data: np.ndarray, times: np.ndarray, r: np.ndarray) -> str:
    requested_t = [100, 300, 600, 900, 1200, 1500, 1800]
    requested_r = [0.0, 0.5, 1.0, 1.5, 2.0]
    t_index = {int(round(t)): i for i, t in enumerate(times)}
    r_index = {round(float(x * 100.0), 10): i for i, x in enumerate(r)}
    lines = [title, "time_s\t" + "\t".join(f"r={x:g}cm" for x in requested_r)]
    for tv in requested_t:
        row = data[t_index[tv]]
        vals = [row[r_index[round(x, 10)]] for x in requested_r]
        lines.append(str(tv) + "\t" + "\t".join(f"{v:.4f}" for v in vals))
    return "\n".join(lines)


def main() -> None:
    # Fine run used for result1; the coarser run is an independent convergence check.
    times, r, temp, moist = solve_q1(dr_cm=0.05, dt=0.5)
    coarse_t, coarse_r, temp_c, moist_c = solve_q1(dr_cm=0.1, dt=1.0)
    fine_idx = [int(round(t - 1.0)) for t in [100, 300, 600, 900, 1200, 1500, 1800]]
    coarse_idx = fine_idx
    # Compare the fine solution sampled at the coarse grid nodes.
    temp_diff = float(np.max(np.abs(temp[fine_idx][:, ::2] - temp_c[coarse_idx])))
    moist_diff = float(np.max(np.abs(moist[fine_idx][:, ::2] - moist_c[coarse_idx])))
    actual_output = write_result1(times, r, temp, moist)

    summary = {
        "solver": "axisymmetric radial finite-volume Crank-Nicolson with implicit Picard moisture update",
        "dr_cm": 0.05,
        "dt_s": 0.5,
        "output_seconds": 1800,
        "radial_nodes_internal": int(len(r)),
        "radial_nodes_output": int(round(2.0 / 0.1)) + 1,
        "max_fine_vs_coarse_temperature": temp_diff,
        "max_fine_vs_coarse_moisture": moist_diff,
        "min_temperature": float(temp.min()),
        "max_temperature": float(temp.max()),
        "min_moisture": float(moist.min()),
        "max_moisture": float(moist.max()),
        "result_file": str(actual_output),
    }
    SUMMARY_FILE.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(format_table("表1 温度 (°C)", temp, times, r))
    print()
    print(format_table("表2 水分浓度 (kg/kg)", moist, times, r))
    print()
    print("RUN_SUMMARY")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
