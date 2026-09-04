"""Build the Chinese mathematical-modeling paper draft as a verified DOCX."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
RESULTS = ROOT / "results"
ASSETS = PAPER / "assets"
OUTPUT = PAPER / "中小微企业信贷决策论文初稿.docx"

INK = "17324D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5F6B76"
RED = "9B1C1C"
TABLE_WIDTH_DXA = 9180

FEATURE_ZH = {
    "sales_signed_amount_asinh": "销项净额规模",
    "sales_positive_amount_log": "销项正额规模",
    "sales_recent12_log": "近12月销项规模",
    "sales_top5_share": "前5大客户占比",
    "sales_hhi": "客户集中度 HHI",
    "sales_monthly_cv_log": "销项月度波动",
    "sales_invoice_count_log": "销项发票数量",
    "sales_void_rate": "销项作废率",
    "sales_active_ratio": "销项活跃月份比",
    "sales_counterparty_count_log": "客户数量",
    "purchase_recent12_log": "近12月进项规模",
    "purchase_positive_amount_log": "进项正额规模",
    "purchase_hhi": "供应商集中度 HHI",
    "purchase_top5_share": "前5大供应商占比",
    "sales_avg_abs_amount_log": "销项平均绝对金额",
    "purchase_void_rate": "进项作废率",
    "purchase_trend_norm": "进项趋势",
}


def _set_run_font(
    run, name: str = "宋体", size: float = 10.5, bold: bool | None = None,
    color: str | None = None, italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size, before, after in [
        ("Heading 1", 15, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if size > 11 else INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "宋体"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def _configure_page(doc: Document) -> None:
    # Named override: academic_a4. The content is a Chinese academic paper,
    # so A4 replaces the narrative-proposal preset's shared Letter geometry.
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.35)
        section.bottom_margin = Cm(2.25)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.2)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run("中小微企业的信贷决策｜论文初稿")
        _set_run_font(run, "微软雅黑", 8.5, color=MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(footer.add_run("—  "), "宋体", 9, color=MUTED)
        _add_field(footer, "PAGE")
        _set_run_font(footer.add_run("  —"), "宋体", 9, color=MUTED)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 80, start: int = 110, bottom: int = 80, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "110")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for idx, text in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = text
        _shade_cell(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                _set_run_font(run, "微软雅黑", 8.5, bold=True, color=INK)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(values):
            cells[idx].text = str(text)
            if row_index % 2 == 1:
                _shade_cell(cells[idx], LIGHT_GRAY)
            for paragraph in cells[idx].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    _set_run_font(run, "宋体", 8.5)
    _set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        _set_run_font(paragraph.add_run(bold_lead), "宋体", 10.5, bold=True)
        text = text[len(bold_lead):]
    _set_run_font(paragraph.add_run(text), "宋体", 10.5)


def _add_equation(doc: Document, equation: str, number: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"{equation}    （{number}）")
    _set_run_font(run, "Cambria Math", 10.5)


def _add_figure(doc: Document, filename: str, caption: str, width: float = 5.8) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(
        str(ASSETS / filename), width=Inches(width)
    )
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    caption_para = doc.add_paragraph(style="Caption")
    _set_run_font(caption_para.add_run(caption), "宋体", 9, color=MUTED)


def _add_section_heading(doc: Document, text: str, level: int = 1, page_break: bool = False) -> None:
    if page_break:
        doc.add_page_break()
    doc.add_heading(text, level=level)


def _cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(85)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(kicker.add_run("2020 年全国大学生数学建模竞赛 C 题"), "微软雅黑", 12, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(12)
    _set_run_font(title.add_run("中小微企业的信贷决策"), "微软雅黑", 28, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(subtitle.add_run("信用风险迁移评价、利率需求响应与多情景鲁棒优化"), "微软雅黑", 14, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(105)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(info.add_run("论文初稿｜可复现建模稿"), "微软雅黑", 11, bold=True, color=BLUE)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(date.add_run("2026 年 9 月"), "宋体", 10.5, color=MUTED)
    doc.add_page_break()


def _front_matter(doc: Document, summary: dict) -> None:
    doc.add_heading("摘  要", level=1)
    abstract = (
        "针对银行在小样本、有标签企业与大样本、无信贷记录企业之间制定年度授信策略的问题，本文建立“经营特征—信用风险—利率响应—组合决策”的递进模型。"
        "首先对企业信息和进销项发票进行语义审计，保留冲销、退款等负数记录，在统一时间窗内构造规模、活跃度、波动、趋势、集中度和异常发票等32维经营特征。"
        "以信誉评级为主监督标签，采用多项 Logistic 与浅层 XGBoost 的折外概率融合构造信用风险指数 CRI；重复分层交叉验证得到融合模型 Accuracy 0.6016、Macro-F1 0.6086、二次加权 Kappa 0.6666。"
        "历史违约只用于检验 CRI 的风险方向及形成评级损失尺度，折外 CRI 对历史违约的 AUC 为0.9066，不将其解释为已校准的企业违约概率。"
        "随后利用保序回归拟合附件3中利率与客户流失率的单调关系，把冻结的风险模型迁移至302家企业，并在单户10万至100万元、年利率4%至15%的范围内求解信贷组合。"
        "在1亿元预算和1.30%组合风险敞口上限下，名义策略选择101家企业，预算误差为0，最小授信23.19万元，期望年度效益142.90万元。"
        "最后由发票波动、近期下滑、上下游集中度、负数发票率和活跃度构造经营脆弱度，设置轻度、中度和重度压力情景，通过 maximin 混合整数规划最大化最坏情景效益。"
        "三种情景下组合效益分别为132.46万元、114.75万元和90.99万元；压力强度提高到1.5倍时最坏效益仍为69.96万元。"
        "SHAP 结果表明，销项净额规模、正额规模、近12月销售和客户集中度是风险评价的主要解释变量。本文同时给出风险上限、LGD和冲击强度敏感性，以区分数据事实与银行风险偏好假设。"
    )
    _add_body(doc, abstract)
    keywords = doc.add_paragraph()
    keywords.paragraph_format.first_line_indent = Cm(0)
    _set_run_font(keywords.add_run("关键词："), "黑体", 10.5, bold=True, color=INK)
    _set_run_font(keywords.add_run("中小微企业；信用风险迁移；XGBoost；保序回归；鲁棒优化；SHAP"), "宋体", 10.5)
    doc.add_page_break()
    doc.add_heading("目  录", level=1)
    contents = [
        "1  问题重述与总体分析", "2  数据审计、模型假设与符号", "3  经营特征与信用风险评价",
        "4  问题一：有信贷记录企业的信贷策略", "5  问题二：风险迁移与一亿元信贷组合",
        "6  问题三：突发事件下的鲁棒信贷策略", "7  模型解释、敏感性与稳健性",
        "8  模型评价、局限与推广", "9  结论", "附录  复现信息与参考文献",
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1.2 if item.startswith(("附",)) else 0.5)
        p.paragraph_format.space_after = Pt(7)
        _set_run_font(p.add_run(item), "宋体", 11)
    doc.add_page_break()


def _section_problem_analysis(doc: Document) -> None:
    doc.add_heading("1  问题重述与总体分析", level=1)
    _add_body(doc, "银行需要依据企业经营信息、历史信誉及客户对贷款利率的响应，在总额度受限时决定“贷给谁、贷多少、以何种利率贷”。附件1提供123家有信贷记录企业，可用于识别经营特征与信誉评级之间的关系；附件2只有302家企业的经营与发票信息，需要把问题一学得的关系迁移过去；附件3给出不同信誉等级客户在若干年利率下的流失率，是利率决策的需求侧依据。")
    _add_body(doc, "三问之间存在明确的因果链。问题一先建立可解释的信用风险评价，并把风险排序转化为固定额度下的授信方案；问题二冻结该评价关系，在无历史信贷标签的企业上估计风险，同时将利率—流失率函数加入组合效益；问题三不重新训练风险模型，而是对风险和客户留存施加参数化压力，寻找在全部情景中最坏收益最大的组合。这样既避免在302家企业上凭空制造标签，也使突发事件只作用于可解释的风险与需求通道。")
    _add_body(doc, "本文不以单一综合评分替代决策。CRI 用于描述信用质量，保序流失曲线描述价格响应，MILP 决定额度组合，SHAP 则用于追溯风险来源。各模块输入输出保持分离，使信用评价误差、风险偏好和压力参数能够分别检查。")


def _section_data(doc: Document) -> None:
    doc.add_heading("2  数据审计、模型假设与符号", level=1)
    doc.add_heading("2.1  数据规模与清洗原则", level=2)
    _add_body(doc, "发票表按照字段语义而不是工作表顺序识别进项与销项。企业与发票的连接键为企业代号；精确重复的8字段整行记录删除，而同一发票号码在不同日期或金额下出现时保留，因为发票号码可能跨年度复用。负数金额代表退款、折让或冲销，若直接取绝对值会破坏净现金流方向，因此保留其符号并另行构造负数比例。")
    _add_table(
        doc,
        ["数据集", "企业数", "进项发票", "销项发票", "负数占比（进/销）"],
        [
            ["附件1", "123", "210,947", "162,484", "0.89% / 5.27%"],
            ["附件2", "302", "395,175", "330,835", "0.92% / 1.72%"],
            ["附件3", "29个利率点", "—", "—", "A/B/C 三条流失曲线"],
        ],
        [1500, 1100, 1700, 1700, 3180],
    )
    _add_body(doc, "金额一致性核验发现5条“金额+税额”与“价税合计”差异超过0.01元，最大差异1.38元，量级远小于企业年度经营规模。模型统一使用来源表中的金额字段，不用价税合计反算金额。全部发票的观察截止日固定为2020年2月21日，避免利用该日之后的信息。")
    doc.add_heading("2.2  标签边界与必要假设", level=2)
    _add_table(
        doc,
        ["假设", "理由与影响"],
        [
            ["信誉评级为主监督标签", "四级评级覆盖全部123家企业，适合做信用风险评价；违约只有27例，不支持复杂个体违约预测。"],
            ["D级企业在问题一禁贷", "附件1中24家D级企业全部违约，且题目要求结合信誉等级制定策略。"],
            ["贷款期限为一年", "与题面一致，效益均按一年名义口径计算，不做跨期贴现。"],
            ["主情景 LGD=1", "附件没有抵押物和回收率信息，采用保守全损口径，并在0.4—1.0范围做敏感性。"],
            ["组合风险敞口上限为1.30%", "题面没有给出银行风险偏好；该值是主决策参数，另给1.28%—1.40%前沿。"],
            ["压力系数不是疫情观测值", "附件缺少行业标签和实际冲击损失，情景只用于压力测试。"],
        ],
        [2300, 6880],
    )
    doc.add_heading("2.3  主要符号", level=2)
    _add_table(
        doc,
        ["符号", "含义", "单位/范围"],
        [
            ["pᵢk", "企业 i 属于评级 k 的融合概率", "[0,1]"],
            ["CRIᵢ", "企业 i 的信用风险指数", "[0,100]"],
            ["ρᵢ", "由评级概率映射的风险损失尺度", "[0,1]"],
            ["ℓᵢ(r)", "利率 r 下企业 i 的客户流失率", "[0,1]"],
            ["xᵢ", "对企业 i 的授信额度", "元"],
            ["vᵢ", "企业 i 的经营脆弱度", "[0,1]"],
            ["uᵢs", "压力情景 s 下单位贷款效益", "元/元"],
        ],
        [1500, 5000, 2680],
    )


def _section_risk_model(doc: Document, summary: dict) -> None:
    doc.add_heading("3  经营特征与信用风险评价", level=1)
    doc.add_heading("3.1  发票经营特征", level=2)
    _add_body(doc, "企业经营状态不能由单一销售额刻画。本文分别在进项和销项方向统计正额、负额、净额、发票数、交易对手数、有效月份数、最近12月规模、线性趋势、增长率、月度变异系数、交易对手 HHI、前五大交易对手占比、负数记录比例和作废率，再加上进销比及经营年限类变量，共形成32维特征。对右偏的规模量取 log(1+x)，对有正负的净额使用 asinh 变换。")
    _add_equation(doc, "HHIᵢ = Σⱼ (aᵢⱼ / Σⱼ aᵢⱼ)²", 1)
    _add_body(doc, "HHI 与前五大客户占比共同描述集中度：前者对全部交易对手加权，后者直观反映头部依赖。月度变异系数用标准差除以绝对均值，均值接近0时通过安全除法返回0并保留异常比例特征，避免数值爆炸。")
    _add_table(
        doc,
        ["特征组", "代表变量", "风险含义"],
        [
            ["经营规模", "正额、净额、发票数、交易对手数", "持续经营能力与现金流基础"],
            ["活跃与近期性", "有效月份比、近12月规模", "是否持续开票及当前业务强度"],
            ["波动与趋势", "月度CV、趋势、近期增长", "经营稳定性和下滑风险"],
            ["集中度", "HHI、前五大交易对手占比", "上下游依赖与单点冲击"],
            ["异常发票", "负数比例、作废率", "退款、冲销及交易异常信号"],
        ],
        [1800, 3400, 3980],
    )
    doc.add_heading("3.2  Logistic/XGBoost 概率融合", level=2)
    _add_body(doc, "多项 Logistic 提供平滑的线性基准，浅层 XGBoost 捕捉规模、波动和集中度之间的非线性交互。所有缺失值填补、标准化和单变量特征选择均封装在交叉验证管线内部，测试折不参与任何预处理参数估计。采用5折、重复5次分层交叉验证得到严格折外概率，再在0到1的离散网格上按多分类对数损失选择 XGBoost 融合权重。")
    _add_equation(doc, "pᵢ = ω pᵢˣᵍᵇ + (1−ω) pᵢˡᵒᵍ,    ω=0.80", 2)
    _add_equation(doc, "CRIᵢ = 100·(0·pᵢA + ⅓pᵢB + ⅔pᵢC + pᵢD)", 3)
    metrics = summary["model_metrics"]
    _add_table(
        doc,
        ["模型", "准确率", "宏平均F1", "序数MAE", "二次Kappa", "对数损失"],
        [
            ["Logistic", f"{metrics['logistic']['accuracy']:.4f}", f"{metrics['logistic']['macro_f1']:.4f}", f"{metrics['logistic']['ordinal_mae']:.4f}", f"{metrics['logistic']['quadratic_kappa']:.4f}", f"{metrics['logistic']['log_loss']:.4f}"],
            ["XGBoost", f"{metrics['xgboost']['accuracy']:.4f}", f"{metrics['xgboost']['macro_f1']:.4f}", f"{metrics['xgboost']['ordinal_mae']:.4f}", f"{metrics['xgboost']['quadratic_kappa']:.4f}", f"{metrics['xgboost']['log_loss']:.4f}"],
            ["融合模型", f"{metrics['ensemble']['accuracy']:.4f}", f"{metrics['ensemble']['macro_f1']:.4f}", f"{metrics['ensemble']['ordinal_mae']:.4f}", f"{metrics['ensemble']['quadratic_kappa']:.4f}", f"{metrics['ensemble']['log_loss']:.4f}"],
        ],
        [1500, 1300, 1400, 1400, 1680, 1900],
    )
    _add_body(doc, "融合模型在 Accuracy、Macro-F1、二次加权 Kappa 和对数损失上均优于两个分量；序数 MAE 略高于 Logistic，说明树模型提高了类别判别但概率期望仍有小幅序数偏差。折外 CRI 在历史违约样本中的均值为75.46，在未违约样本中为40.67，AUC 为0.9066，验证了指数的风险方向。该 AUC 受 D 级与违约高度相关影响，只作为外部一致性证据。")
    _add_figure(doc, "risk_index_distribution.png", "图1  附件1折外 CRI 与附件2迁移 CRI 的分布")


def _section_q1(doc: Document, p1: pd.DataFrame, curve1: pd.DataFrame) -> None:
    doc.add_heading("4  问题一：有信贷记录企业的信贷策略", level=1)
    doc.add_heading("4.1  风险损失尺度与利率响应", level=2)
    _add_body(doc, "各评级的历史违约样本数差异较大，直接用频率会使 A 级违约率变成0。本文采用 Jeffreys 平滑，将评级 k 的违约损失尺度写为下式，其中 dₖ、nₖ分别为违约数与企业数。A/B/C/D 四档得到0.0179、0.0385、0.0714和0.9800。")
    _add_equation(doc, "d̃ₖ = (dₖ + 0.5) / (nₖ + 1),    ρᵢ = Σₖ pᵢk d̃ₖ", 4)
    _add_body(doc, "附件3的原始流失率存在少量局部下降，直接插值会意味着利率上升反而减少流失。故分别对 A、B、C 级应用保序回归，求与原数据平方误差最小的非降序列；企业流失率再按其 A/B/C 概率加权。D 类未给流失曲线，在问题一直接禁贷，在问题二通过高风险损失尺度约束。")
    _add_figure(doc, "interest_loss_isotonic.png", "图2  年利率—客户流失率原始点与保序拟合")
    _add_body(doc, "设 r 为年利率、LGD 为违约损失率，客户接受贷款的概率为1−ℓᵢ(r)。单位授信的期望效益由名义利息扣除客户流失造成的利息损失和风险损失：")
    _add_equation(doc, "cᵢ(r) = [r − LGD·ρᵢ]·[1−ℓᵢ(r)]", 5)
    doc.add_heading("4.2  固定额度 MILP", level=2)
    _add_body(doc, "对每家企业先在附件3的29个利率点中选择使 cᵢ(r) 最大的利率，再在固定预算 B 下配置额度。令 yᵢ 表示是否授信，额度上下限分别为10万元和100万元，D 级企业的 yᵢ 固定为0。")
    _add_equation(doc, "max Σᵢ cᵢ xᵢ", 6)
    _add_equation(doc, "s.t. Σᵢxᵢ=B;  10yᵢ≤xᵢ≤100yᵢ;  Σᵢρᵢ(1−ℓᵢ)xᵢ≤R;  yᵢ∈{0,1}", 7)
    q1_groups = p1.groupby("rating_observed").agg(企业数=("cri_oof", "size"), 平均CRI=("cri_oof", "mean"), 入选数=("selected", "sum"), 授信额=("loan_yuan", "sum"))
    rows = [[idx, str(int(row["企业数"])), f"{row['平均CRI']:.2f}", str(int(row["入选数"])), f"{row['授信额']/1e4:.0f}"] for idx, row in q1_groups.iterrows()]
    _add_table(doc, ["评级", "企业数", "平均CRI", "入选数", "授信额/万元"], rows, [1300, 1500, 1800, 1500, 3080])
    _add_body(doc, "题面只说明问题一年度总额度固定，没有给出具体数值。本文取5000万元作为展示主情景，并以2000万元和8000万元检验预算变化。主情景选择50家企业，其中 A/B/C 级分别为21、16和13家，D 级为0；由于单位效益在额度上是线性的，优先企业达到100万元上限，这是线性信贷收益模型的角点结果。")
    _add_figure(doc, "problem1_budget_curve.png", "图3  问题一不同展示预算下的组合期望效益")


def _section_q2(doc: Document, p2: pd.DataFrame, curve2: pd.DataFrame, shift: pd.DataFrame) -> None:
    doc.add_heading("5  问题二：风险迁移与一亿元信贷组合", level=1)
    doc.add_heading("5.1  冻结模型迁移与漂移检查", level=2)
    _add_body(doc, "问题二不重新拟合分类器，而是将问题一的全样本模型、14个经训练管线选择的变量、融合权重和评级顺序全部冻结，再对附件2的32维同构特征输出评级概率与 CRI。这是有监督关系的直接迁移，而不是把302家企业重新聚类后人为命名等级。")
    top_shift = shift.head(6)
    _add_table(
        doc,
        ["特征", "源域均值", "目标域均值", "标准化均值漂移"],
        [[FEATURE_ZH.get(r.feature, r.feature), f"{r.source_mean:.3f}", f"{r.target_mean:.3f}", f"{r.standardized_mean_shift:+.3f}"] for r in top_shift.itertuples()],
        [2800, 1800, 1800, 2780],
    )
    _add_body(doc, "绝对漂移最大的变量为进项 HHI，标准化均值差仅0.219；其余前列变量也均小于0.18。按企业32维标准化偏差的平均绝对值计算迁移置信度，302家企业均未超过1.5的人工复核阈值。该结果说明均值层面的域差异较小，但目标域缺少真实评级，因此不能据此宣称迁移准确率已经验证。")
    doc.add_heading("5.2  一亿元主组合与风险前沿", level=2)
    _add_body(doc, "将迁移评级概率映射为风险损失尺度，并按照式（5）选择利率。问题二允许所有企业进入候选集，但预测为 D 型的企业会因较高 ρᵢ 受到目标函数和1.30%组合风险上限的连续惩罚。额度仍限制在10万至100万元，总额为1亿元。")
    rows = []
    for row in curve2.itertuples():
        rows.append([
            f"{100*row.risk_cap_fraction:.2f}%", f"{row.expected_benefit_yuan/1e4:.2f}",
            f"{row.risk_loss_principal_yuan/1e4:.2f}", str(int(row.selected_count)),
            f"{row.min_selected_loan_yuan/1e4:.2f}",
        ])
    _add_table(doc, ["风险上限", "期望效益/万元", "风险尺度/万元", "入选数", "最小额度/万元"], rows, [1500, 2100, 2100, 1500, 1980])
    main_selected = p2[p2["selected"]]
    rating_counts = main_selected["rating_predicted"].value_counts().reindex(["A", "B", "C", "D"], fill_value=0)
    partial = main_selected[main_selected["loan_yuan"] < 999_999.0].sort_values("loan_yuan")
    _add_body(doc, f"主情景在1.30%风险上限下选择101家企业，预测评级 A/B/C/D 的入选数分别为{rating_counts['A']}、{rating_counts['B']}、{rating_counts['C']}和{rating_counts['D']}家；总额度严格等于1亿元，期望效益142.90万元，风险损失本金尺度130.00万元。最小授信为23.19万元，最大为100万元。共有{len(partial)}家企业未达到100万元上限，说明风险约束在边界企业上产生了额度调整。")
    _add_figure(doc, "problem2_risk_return_frontier.png", "图4  问题二风险上限变化下的组合风险—效益前沿")
    _add_body(doc, "风险上限从1.28%提高至1.34%时，期望效益由140.02万元增至144.91万元；超过1.34%后曲线变平，表明风险约束不再是活跃约束。选择1.30%作为主情景可牺牲约2.01万元名义效益，换取约3.37万元风险敞口下降，并将授信从100家扩展到101家。")


def _section_q3(doc: Document, scenarios: pd.DataFrame, stress: pd.DataFrame) -> None:
    scenario_zh = {"mild": "轻度", "moderate": "中度", "severe": "重度"}
    doc.add_heading("6  问题三：突发事件下的鲁棒信贷策略", level=1)
    doc.add_heading("6.1  发票经营脆弱度", level=2)
    _add_body(doc, "附件缺少行业类别、地区、疫情暴露和实际冲击损失。若直接套用外部行业系数，会产生无法核验的企业级影响。本文只用附件内可观察的经营韧性构造脆弱度：销项月度波动、销售近期下滑、客户 HHI、供应商 HHI、销项负数率和低活跃度分别做百分位排序后等权平均。")
    _add_equation(doc, "vᵢ = [P(CVᵢ)+1−P(gᵢ)+P(HHIˢᵢ)+P(HHIᵖᵢ)+P(negᵢ)+1−P(activeᵢ)]/6", 8)
    doc.add_heading("6.2  压力情景与 maximin 模型", level=2)
    _add_body(doc, "设情景强度参数 δₛ 分别为0.15、0.50、1.00，客户留存冲击 γₛ 分别为0.05、0.15、0.30。风险损失尺度按脆弱度相对放大，留存率按脆弱度下降：")
    _add_equation(doc, "ρᵢs = min{1, ρᵢ(1+δₛvᵢ)},    qᵢs(r)=[1−ℓᵢ(r)](1−γₛvᵢ)", 9)
    _add_equation(doc, "uᵢs(r) = [r−LGD·ρᵢs] qᵢs(r)", 10)
    _add_body(doc, "对每家企业在29个利率点中选择最大化最坏单位效益的利率；组合层引入辅助变量 z，使 z 不超过任一情景下的总效益。")
    _add_equation(doc, "max z,    s.t. z≤Σᵢuᵢs xᵢ  (∀s),  Σᵢxᵢ=10⁸,  Σᵢmaxₛ{ρᵢs qᵢs}xᵢ≤1.3×10⁶", 11)
    _add_table(
        doc,
        ["情景", "组合效益/万元", "风险尺度/万元", "最坏效益约束"],
        [[scenario_zh[str(r.scenario)], f"{r.portfolio_benefit_yuan/1e4:.2f}", f"{r.portfolio_risk_yuan/1e4:.2f}", "活跃" if r.scenario == "severe" else "非活跃"] for r in scenarios.itertuples()],
        [1700, 2500, 2500, 2480],
    )
    _add_body(doc, "鲁棒策略在轻度、中度和重度压力下的组合效益分别为132.46万元、114.75万元和90.99万元，重度情景构成 maximin 的活跃约束。三种情景的风险损失尺度均未超过130万元；主鲁棒组合选择100家企业并使用全部1亿元预算。由于严重压力同时提高违约损失和客户流失，稳健利率的贷款加权均值提高到13.554%，以利差补偿风险，但仍受15%上限约束。")
    _add_figure(doc, "problem3_robust_scenarios.png", "图5  鲁棒组合在轻度、中度与重度压力下的年度效益")


def _section_explain(doc: Document, shap: pd.DataFrame, lgd: pd.DataFrame, stress: pd.DataFrame) -> None:
    doc.add_heading("7  模型解释、敏感性与稳健性", level=1)
    doc.add_heading("7.1  SHAP 全局与局部解释", level=2)
    _add_body(doc, "TreeSHAP 只解释融合模型中的 XGBoost 分量。对四类 SHAP 值按照 CRI 的0、1/3、2/3、1严重度方向加权，可得到每个特征对“风险方向”的贡献。全局重要性取该方向贡献的平均绝对值；局部解释则给出高风险企业中推动 CRI 上升的前三个变量。")
    top = shap.head(10)
    _add_table(
        doc,
        ["排名", "特征", "平均绝对风险贡献", "平均方向贡献"],
        [[str(i), FEATURE_ZH.get(r.feature, r.feature), f"{r.mean_abs_shap_cri_direction:.4f}", f"{r.mean_shap_cri_direction:+.4f}"] for i, r in enumerate(top.itertuples(), start=1)],
        [1000, 3600, 2500, 2080],
    )
    _add_figure(doc, "shap_global_importance.png", "图6  XGBoost 分量的全局 SHAP 风险方向重要性")
    _add_body(doc, "销项净额规模、销项正额规模和近12月销项规模占据前三位，说明模型首先使用销售规模与持续性区分风险；前五大客户占比和客户 HHI 紧随其后，反映对少数客户的依赖。平均方向贡献不能单独用于判断“数值大必然增险”，因为树模型存在阈值和交互；对单个企业应结合其实际特征值与局部 SHAP 正负号解释。")
    doc.add_heading("7.2  LGD 与压力强度敏感性", level=2)
    _add_body(doc, "LGD 敏感性用于观察风险损失估值变化，诊断中放宽硬风险上限以隔离 LGD 对利率和效益的影响。LGD 从0.4提高到1.0时，组合期望效益从265.71万元降至144.91万元，贷款加权利率从7.834%升至11.994%。较高 LGD 使模型偏向更高利率和更低留存的组合，以控制被接受贷款上的损失敞口。")
    _add_table(
        doc,
        ["LGD", "期望效益/万元", "风险尺度/万元", "贷款加权利率"],
        [[f"{r.lgd:.1f}", f"{r.expected_benefit_yuan/1e4:.2f}", f"{r.risk_loss_principal_yuan/1e4:.2f}", f"{100*r.loan_weighted_rate:.3f}%"] for r in lgd.itertuples()],
        [1600, 2600, 2600, 2380],
    )
    _add_body(doc, "将三档情景的全部 δₛ、γₛ 同时乘以0.5至1.5，最坏效益由115.38万元单调下降至69.96万元，风险尺度始终小于130万元；稳健贷款加权利率从13.162%升至13.890%。这说明主结论不是单一冲击点的偶然产物，但利润水平对冲击强度具有明显敏感性。")
    _add_figure(doc, "sensitivity_analysis.png", "图7  LGD 与压力强度敏感性分析")


def _section_evaluation(doc: Document) -> None:
    doc.add_heading("8  模型评价、局限与推广", level=1)
    doc.add_heading("8.1  模型优点", level=2)
    _add_body(doc, "第一，标签语义清晰。模型学习的是信誉评级而不是用27个违约样本强行拟合高维违约概率；所有分类指标来自严格折外预测。第二，风险、需求和组合决策分层，避免将客户流失混入信用评分。第三，附件3的单调性由保序回归保证，既保留数据形状，又消除局部违反经济常识的波动。第四，突发事件模型只基于发票可观察韧性，情景参数和数据事实明确分开。第五，CSV、运行日志、随机种子和产物哈希均被保存，预算、额度、利率、D级禁贷和求解状态有自动校验。")
    doc.add_heading("8.2  局限性", level=2)
    _add_body(doc, "源域只有123家企业，评级类别之间的经营模式可能重叠，融合模型 Macro-F1 约0.61，说明 CRI 适合排序和组合约束，不适合替代人工授信审查。目标域没有真实评级，均值漂移较小只能支持“可迁移性没有明显反证”，不能验证个体预测准确率。评级到风险损失尺度的映射受到 D 级全部违约这一样本结构影响，也缺少抵押物、回收率、财务报表和企业行业信息。")
    _add_body(doc, "额度收益按贷款金额线性计算，因此在风险约束不活跃时会出现若干企业达到100万元上限的角点。若银行希望主动分散集中风险，可在后续模型中加入企业/行业集中度约束或分段递减的边际效益。当前利率先按企业单位效益选择，再进行组合分配；在极紧风险上限下，可进一步建立企业—利率联合整数变量，使利率选择与组合风险约束完全联立。")
    doc.add_heading("8.3  推广建议", level=2)
    _add_body(doc, "实际部署时可按月滚动更新发票特征，保留模型评分、SHAP 原因和人工审批结果。当积累到足够的真实违约、回收和拒贷数据后，应分别校准 PD、LGD 与客户接受概率，并用时间外验证替代随机交叉验证。若补充行业标签，可采用行业层次模型和行业额度上限；若获得宏观情景预测，可把本文的参数化压力集替换为带概率或区间的场景树。")


def _section_conclusion(doc: Document) -> None:
    doc.add_heading("9  结论", level=1)
    _add_body(doc, "本文以发票经营行为为统一信息载体，完成了从有信贷记录企业的信用评价，到无信贷记录企业的风险迁移，再到突发事件下鲁棒调整的完整决策链。问题一的 Logistic/XGBoost 融合模型以80%权重吸收树模型的非线性能力，折外 Macro-F1 为0.6086、二次加权 Kappa 为0.6666；在5000万元展示预算下选择50家 A/B/C 级企业，未向 D 级企业授信。")
    _add_body(doc, "问题二冻结问题一模型迁移到302家企业，并将保序利率需求响应与风险损失尺度共同纳入组合效益。在1亿元预算、单户10万至100万元和1.30%风险上限下，主策略选择101家，最小授信23.19万元，预算误差为0，期望年度效益142.90万元。问题三使用经营脆弱度构造三档压力集，maximin 策略在重度情景下仍保持90.99万元效益；压力提高到1.5倍时最坏效益为69.96万元。")
    _add_body(doc, "上述结果表明，信用评分只负责度量风险，利率响应负责描述客户选择，组合优化负责落实银行风险偏好，三者分离后更容易解释和更新。对实际银行而言，1.30%风险上限、LGD 和压力系数应由风险管理部门结合资本成本与历史回收率重新标定；本文给出的策略是可复现的决策基线，而不是未经校准的监管违约概率。")


def _appendix(doc: Document, summary: dict) -> None:
    doc.add_page_break()
    doc.add_heading("附录A  主要参数与复现信息", level=1)
    _add_table(
        doc,
        ["项目", "取值"],
        [
            ["随机种子", str(summary["seed"])],
            ["交叉验证", "5折×5次重复分层"],
            ["XGBoost", "160棵树，深度2，学习率0.04"],
            ["融合权重", "XGBoost 0.80，Logistic 0.20"],
            ["贷款边界", "10万—100万元，期限1年"],
            ["利率网格", "附件3的29个利率点，4%—15%"],
            ["问题一展示预算", "5000万元；另算2000万元、8000万元"],
            ["问题二/三预算", "1亿元"],
            ["主风险敞口上限", "预算的1.30%"],
            ["主 LGD", "1.0"],
        ],
        [3100, 6080],
    )
    _add_body(doc, "复现环境为 Windows 11、Python 3.12.13。关键软件版本包括 NumPy 2.5.2、pandas 3.0.5、scikit-learn 1.9.0、SciPy 1.18.1、XGBoost 3.4.1、SHAP 0.52.0 和 matplotlib 3.11.1。模型主入口、6项测试和代码审计均在项目隔离环境中成功运行。")
    doc.add_heading("附录B  参考文献", level=1)
    references = [
        "[1] 全国大学生数学建模竞赛组委会. 2020年全国大学生数学建模竞赛C题：中小微企业的信贷决策, 2020.",
        "[2] Chen T, Guestrin C. XGBoost: A Scalable Tree Boosting System. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 2016: 785-794. DOI: 10.1145/2939672.2939785.",
        "[3] Lundberg S M, Lee S I. A Unified Approach to Interpreting Model Predictions. Advances in Neural Information Processing Systems 30, 2017.",
        "[4] Barlow R E, Bartholomew D J, Bremner J M, Brunk H D. Statistical Inference Under Order Restrictions. Wiley, 1972.",
        "[5] Ben-David S, Blitzer J, Crammer K, et al. A Theory of Learning from Different Domains. Machine Learning, 2010, 79: 151-175. DOI: 10.1007/s10994-009-5152-4.",
        "[6] Bertsimas D, Sim M. The Price of Robustness. Operations Research, 2004, 52(1): 35-53. DOI: 10.1287/opre.1030.0065.",
        "[7] Hosmer D W, Lemeshow S, Sturdivant R X. Applied Logistic Regression. 3rd ed. Wiley, 2013.",
        "[8] Agresti A. Categorical Data Analysis. 2nd ed. Wiley, 2002.",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.space_after = Pt(4)
        _set_run_font(p.add_run(ref), "宋体", 9.5)


def build() -> Path:
    summary = json.loads((RESULTS / "summary.json").read_text(encoding="utf-8"))
    p1 = pd.read_csv(RESULTS / "problem1_strategy.csv")
    p2 = pd.read_csv(RESULTS / "problem2_strategy.csv")
    curve1 = pd.read_csv(RESULTS / "problem1_budget_curve.csv")
    curve2 = pd.read_csv(RESULTS / "problem2_risk_return_curve.csv")
    scenarios = pd.read_csv(RESULTS / "problem3_scenarios.csv")
    stress = pd.read_csv(RESULTS / "problem3_stress_sensitivity.csv")
    lgd = pd.read_csv(RESULTS / "problem2_lgd_sensitivity.csv")
    shap = pd.read_csv(RESULTS / "shap_global.csv")
    shift = pd.read_csv(RESULTS / "transfer_feature_shift.csv")

    doc = Document()
    _configure_styles(doc)
    _configure_page(doc)
    doc.core_properties.title = "中小微企业的信贷决策"
    doc.core_properties.subject = "信用风险迁移评价、利率需求响应与多情景鲁棒优化"
    doc.core_properties.author = "数学建模团队"
    doc.core_properties.keywords = "信用风险迁移, XGBoost, 鲁棒优化, SHAP"

    _cover(doc)
    _front_matter(doc, summary)
    _section_problem_analysis(doc)
    _section_data(doc)
    _section_risk_model(doc, summary)
    _section_q1(doc, p1, curve1)
    _section_q2(doc, p2, curve2, shift)
    _section_q3(doc, scenarios, stress)
    _section_explain(doc, shap, lgd, stress)
    _section_evaluation(doc)
    _section_conclusion(doc)
    _appendix(doc, summary)

    # Update fields on open; Word/LibreOffice may refresh page fields during export.
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
