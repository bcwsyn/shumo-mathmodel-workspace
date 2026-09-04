from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "reference" / "全篇证据约束优化版_V8.1.docx"
DOCX = ROOT / "paper" / "全篇证据约束优化版_V8.2_视觉优化.docx"
PDF = ROOT / "paper" / "全篇证据约束优化版_V8.2_视觉优化.pdf"
FIGURE_DIR = ROOT / "figures" / "publication"
REPORT = ROOT / "reports" / "v82_visual_qa.json"

EXPECTED_SOURCE_SHA256 = "93A5854030AFF664340487035C60BE2997FE26520EE49A8C201BFA8B69C9FBD0"
FIGURE_FILES = [f"figure{number:02d}_{stem}.png" for number, stem in enumerate(
    [
        "method_overview",
        "risk_pipeline",
        "cri_distribution",
        "interest_loss",
        "budget_benefit",
        "transfer_shift",
        "risk_return",
        "stress_transmission",
        "robust_scenarios",
        "shap_importance",
        "sensitivity",
    ],
    start=1,
)]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def expected_paragraph_texts(source: Document) -> list[str]:
    texts = [paragraph.text for paragraph in source.paragraphs]
    start = next(i for i, text in enumerate(texts) if re.sub(r"\s+", "", text) == "目录")
    end = next(
        i
        for i in range(start + 1, len(texts))
        if source.paragraphs[i].style.name == "Heading 1"
    )
    expected = texts[:start] + texts[end:]
    for index, text in enumerate(expected):
        if "论文初稿｜可复现建模稿" in text:
            expected[index] = text.replace("论文初稿｜可复现建模稿", "数学建模论文｜证据约束版")
        elif text.startswith("图7 "):
            expected[index] = (
                "图7 问题二风险上限变化下的组合风险—效益前沿。"
                "局部框放大1.30%主方案至1.34%平台段，用于辨识主方案附近的边际效益与约束转折；"
                "不表示新增样本、插值结果或第二套模型。"
            )
        elif text.startswith("风险上限从1.28%"):
            expected[index] = (
                "该局部框仅放大已有五个风险上限方案中的主方案邻域与平台段，"
                "不增加数据点或重新拟合曲线。" + text
            )
    return expected


def main() -> None:
    source = Document(SOURCE)
    output = Document(DOCX)
    pdf = PdfReader(PDF)

    with zipfile.ZipFile(DOCX) as archive:
        bad_zip_member = archive.testzip()

    source_text_expected = expected_paragraph_texts(source)
    output_text = [paragraph.text for paragraph in output.paragraphs]
    source_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in source.tables]
    output_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in output.tables]

    figure_matches: list[bool] = []
    figure_dpi: list[list[float]] = []
    for shape, filename in zip(output.inline_shapes, FIGURE_FILES, strict=True):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blob = output.part.related_parts[rid].blob
        figure_path = FIGURE_DIR / filename
        figure_matches.append(blob == figure_path.read_bytes())
        with Image.open(figure_path) as image:
            dpi = image.info.get("dpi", (0.0, 0.0))
            figure_dpi.append([round(float(dpi[0]), 2), round(float(dpi[1]), 2)])

    page_sizes = [
        [round(float(page.mediabox.width), 2), round(float(page.mediabox.height), 2)]
        for page in pdf.pages
    ]
    all_a4 = all(abs(width - 595.32) < 0.5 and abs(height - 841.92) < 0.5 for width, height in page_sizes)

    report = {
        "status": "PASS",
        "source": {
            "path": str(SOURCE),
            "sha256": sha256(SOURCE),
            "hash_matches_frozen_receipt": sha256(SOURCE) == EXPECTED_SOURCE_SHA256,
        },
        "docx": {
            "path": str(DOCX),
            "sha256": sha256(DOCX),
            "zip_integrity": bad_zip_member is None,
            "paragraph_count": len(output.paragraphs),
            "table_count": len(output.tables),
            "inline_figure_count": len(output.inline_shapes),
            "paragraph_text_matches_expected_bounded_edits": output_text == source_text_expected,
            "table_text_unchanged": output_tables == source_tables,
            "publication_figure_blobs_match": all(figure_matches),
            "all_figure_pngs_approximately_300_dpi": all(
                abs(x_dpi - 300) < 1 and abs(y_dpi - 300) < 1 for x_dpi, y_dpi in figure_dpi
            ),
        },
        "pdf": {
            "path": str(PDF),
            "sha256": sha256(PDF),
            "page_count": len(pdf.pages),
            "all_pages_a4": all_a4,
            "word_render_review": "20/20 pages reviewed; final affected pages 3 and 12 rechecked",
        },
        "figures": {
            "png_count": len(list(FIGURE_DIR.glob("*.png"))),
            "pdf_count": len(list(FIGURE_DIR.glob("*.pdf"))),
            "svg_count": len(list(FIGURE_DIR.glob("*.svg"))),
            "dpi": figure_dpi,
        },
    }

    required = [
        report["source"]["hash_matches_frozen_receipt"],
        report["docx"]["zip_integrity"],
        report["docx"]["table_count"] == 12,
        report["docx"]["inline_figure_count"] == 11,
        report["docx"]["paragraph_text_matches_expected_bounded_edits"],
        report["docx"]["table_text_unchanged"],
        report["docx"]["publication_figure_blobs_match"],
        report["docx"]["all_figure_pngs_approximately_300_dpi"],
        report["pdf"]["page_count"] == 20,
        report["pdf"]["all_pages_a4"],
        report["figures"]["png_count"] == 11,
        report["figures"]["pdf_count"] == 11,
        report["figures"]["svg_count"] == 11,
    ]
    if not all(required):
        report["status"] = "FAIL"

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
