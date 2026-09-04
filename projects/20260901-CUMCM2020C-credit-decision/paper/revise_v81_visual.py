from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "reference" / "全篇证据约束优化版_V8.1.docx"
OUTPUT = ROOT / "paper" / "全篇证据约束优化版_V8.2_视觉优化.docx"
FIGURE_DIR = ROOT / "figures" / "publication"

NAVY = "17324D"
BLUE = "2878B5"
TEXT = "263442"
MUTED = "536273"
GRID = "AAB8C6"
HEADER_FILL = "EAF0F6"
ALT_FILL = "F7F9FB"

FIGURES = [
    ("figure01_method_overview.png", 6.30, "图1 三问统一建模与证据闭环"),
    ("figure02_risk_pipeline.png", 6.10, "图2 信用风险评价与迁移流程"),
    ("figure03_cri_distribution.png", 5.80, "图3 信用风险指数分布"),
    ("figure04_interest_loss.png", 5.80, "图4 贷款年利率与客户流失率关系"),
    ("figure05_budget_benefit.png", 5.80, "图5 授信预算与组合效益关系"),
    ("figure06_transfer_shift.png", 6.10, "图6 有信贷记录与无信贷记录企业的风险分布迁移"),
    ("figure07_risk_return.png", 5.80, "图7 风险上限变化下的组合风险—效益前沿"),
    ("figure08_stress_transmission.png", 6.10, "图8 疫情冲击在模型中的传导路径"),
    ("figure09_robust_scenarios.png", 5.80, "图9 不同情景下的组合收益与损失"),
    ("figure10_shap_importance.png", 5.80, "图10 信用风险模型的特征贡献"),
    ("figure11_sensitivity.png", 5.80, "图11 关键参数敏感性分析"),
]


def set_run_font(run, east_asia: str, size: float, *, color: str | None = None, bold: bool | None = None) -> None:
    run.font.name = east_asia
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_style_font(style, east_asia: str, size: float, *, color: str, bold: bool = False) -> None:
    style.font.name = east_asia
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def delete_manual_toc(doc: Document) -> None:
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if re.sub(r"\s+", "", p.text) == "目录")
    end = next(
        i
        for i in range(start + 1, len(paragraphs))
        if paragraphs[i].style.name == "Heading 1"
    )
    for paragraph in paragraphs[start:end]:
        element = paragraph._element
        element.getparent().remove(element)

    first_heading = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Heading 1" and re.match(r"^1\s+问题重述", paragraph.text.strip())
    )
    replacement = first_heading.insert_paragraph_before(first_heading.text, style="Heading 1")
    replacement.paragraph_format.keep_with_next = True
    element = first_heading._element
    element.getparent().remove(element)


def set_cell_shading(cell, fill: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


def set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def set_row_flag(row, tag_name: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tag = qn(f"w:{tag_name}")
    if tr_pr.find(tag) is None:
        tr_pr.append(OxmlElement(f"w:{tag_name}"))


def style_tables(doc: Document) -> None:
    for table_index, table in enumerate(doc.tables):
        set_table_borders(table)
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            set_row_flag(row, "cantSplit")
            if row_index == 0:
                set_row_flag(row, "tblHeader")
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                compact = table_index == 2
                set_cell_margins(
                    cell,
                    top=34 if compact else 56,
                    start=62 if compact else 78,
                    bottom=34 if compact else 56,
                    end=62 if compact else 78,
                )
                set_cell_shading(cell, HEADER_FILL if row_index == 0 else (ALT_FILL if row_index % 2 == 0 else None))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.keep_together = True
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            "宋体",
                            7.8 if compact else 8.7,
                            color=TEXT,
                            bold=True if row_index == 0 else run.bold,
                        )


def paragraph_has_drawing(paragraph) -> bool:
    return bool(paragraph._element.xpath(".//w:drawing"))


def paragraph_has_equation(paragraph) -> bool:
    return bool(paragraph._element.xpath(".//m:oMath | .//m:oMathPara"))


def normalize_paragraphs(doc: Document) -> None:
    in_references = False
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name

        if text == "附录B 参考文献":
            in_references = True

        if style_name.startswith("Heading 1"):
            paragraph.paragraph_format.space_before = Pt(14)
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.first_line_indent = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, "微软雅黑", 15.5, color=NAVY, bold=True)
            continue

        if style_name.startswith("Heading 2"):
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.first_line_indent = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, "微软雅黑", 12.5, color=BLUE, bold=True)
            continue

        if paragraph_has_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            continue

        if re.match(r"^[图表]\s*\d+\s+", text):
            paragraph.style = doc.styles["Caption"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = text.startswith("表")
            for run in paragraph.runs:
                set_run_font(run, "宋体", 9.0, color=MUTED)
            continue

        if paragraph_has_equation(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True
            continue

        if index < 12:
            continue

        if text.startswith("复现材料索引："):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(5)
            for run in paragraph.runs:
                set_run_font(run, "等线", 8.5, color=TEXT)
            continue

        if in_references and text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0.63)
            paragraph.paragraph_format.first_line_indent = Cm(-0.63)
            paragraph.paragraph_format.line_spacing = 1.2
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                set_run_font(run, "宋体", 9.5, color=TEXT)
            continue

        if text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Pt(21)
            paragraph.paragraph_format.line_spacing = 1.45
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.widow_control = True
            for run in paragraph.runs:
                set_run_font(run, "宋体", 10.5, color=TEXT)


def replace_figures(doc: Document) -> None:
    shapes = list(doc.inline_shapes)
    if len(shapes) != len(FIGURES):
        raise RuntimeError(f"Expected {len(FIGURES)} inline figures, found {len(shapes)}")

    from PIL import Image

    for shape, (filename, display_width, alt_text) in zip(shapes, FIGURES, strict=True):
        image_path = FIGURE_DIR / filename
        with Image.open(image_path) as image:
            pixel_width, pixel_height = image.size
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = doc.part.related_parts[rid]
        image_part._blob = image_path.read_bytes()
        shape.width = Inches(display_width)
        shape.height = Inches(display_width * pixel_height / pixel_width)
        doc_pr = shape._inline.docPr
        doc_pr.set("name", alt_text)
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", alt_text)


def update_text_and_headers(doc: Document) -> None:
    for paragraph in doc.paragraphs[:12]:
        if "论文初稿｜可复现建模稿" in paragraph.text:
            for run in paragraph.runs:
                if "论文初稿｜可复现建模稿" in run.text:
                    run.text = run.text.replace("论文初稿｜可复现建模稿", "数学建模论文｜证据约束版")

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("图7 "):
            paragraph.text = (
                "图7 问题二风险上限变化下的组合风险—效益前沿。"
                "局部框放大1.30%主方案至1.34%平台段，用于辨识主方案附近的边际效益与约束转折；"
                "不表示新增样本、插值结果或第二套模型。"
            )
        elif paragraph.text.startswith("风险上限从1.28%"):
            paragraph.text = (
                "该局部框仅放大已有五个风险上限方案中的主方案邻域与平台段，"
                "不增加数据点或重新拟合曲线。" + paragraph.text
            )

    seen_headers: set[int] = set()
    seen_footers: set[int] = set()
    for section in doc.sections:
        section.different_first_page_header_footer = True
        first_header = section.first_page_header
        first_footer = section.first_page_footer
        for paragraph in first_header.paragraphs:
            paragraph.text = ""
        for paragraph in first_footer.paragraphs:
            paragraph.text = ""

        for header in (section.header, section.even_page_header):
            if id(header._element) in seen_headers:
                continue
            seen_headers.add(id(header._element))
            for paragraph in header.paragraphs:
                if paragraph.text.strip():
                    paragraph.text = "中小微企业的信贷决策"
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in paragraph.runs:
                        set_run_font(run, "微软雅黑", 8.5, color=MUTED)
        for footer in (section.footer, section.even_page_footer):
            if id(footer._element) in seen_footers:
                continue
            seen_footers.add(id(footer._element))
            for paragraph in footer.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, "等线", 8.5, color=MUTED)


def configure_styles(doc: Document) -> None:
    set_style_font(doc.styles["Normal"], "宋体", 10.5, color=TEXT)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.45
    normal.space_after = Pt(3)

    set_style_font(doc.styles["Heading 1"], "微软雅黑", 15.5, color=NAVY, bold=True)
    set_style_font(doc.styles["Heading 2"], "微软雅黑", 12.5, color=BLUE, bold=True)
    set_style_font(doc.styles["Caption"], "宋体", 9.0, color=MUTED)


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    missing = [str(FIGURE_DIR / name) for name, _, _ in FIGURES if not (FIGURE_DIR / name).exists()]
    if missing:
        raise FileNotFoundError("Missing publication figures:\n" + "\n".join(missing))

    doc = Document(SOURCE)
    delete_manual_toc(doc)
    update_text_and_headers(doc)
    replace_figures(doc)
    configure_styles(doc)
    normalize_paragraphs(doc)
    style_tables(doc)
    enable_field_updates(doc)

    props = doc.core_properties
    props.title = "中小微企业的信贷决策——证据约束视觉优化版"
    props.subject = "2020年全国大学生数学建模竞赛C题"
    props.comments = "V8.2：基于V8.1的版式与图形表达优化；模型数据与结论保持不变。"
    props.modified = datetime.now(timezone.utc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
